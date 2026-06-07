"""
tests/test_excel_writeback.py
CPOI Platform -- Tests for intake/excel_writeback.py (Module 7 outbound).

Creates a minimal real .xlsx fixture in a temp directory, runs
write_variance_flags(), and asserts the output workbook structure.
"""

import sys
from dataclasses import dataclass
from pathlib import Path

import pytest
from openpyxl import load_workbook, Workbook

sys.path.insert(0, str(Path(__file__).parent.parent))

from intake.excel_writeback import write_variance_flags, _SHEET_NAME, _HEADERS


# ------------------------------------------------------------------
# Minimal AnomalyFlag stand-in (mirrors scoring.anomaly.AnomalyFlag)
# ------------------------------------------------------------------

@dataclass
class _Flag:
    flag_id: str
    flag_type: str
    dimension: str | None
    description: str
    detected_at: str


# ------------------------------------------------------------------
# Fixtures
# ------------------------------------------------------------------

@pytest.fixture
def source_workbook(tmp_path: Path) -> Path:
    """Create a minimal .xlsx with one sheet of dummy data."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Client_Data"
    ws.append(["Metric", "Value"])
    ws.append(["headcount", 42])
    path = tmp_path / "client_alpha.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def two_flags() -> list[_Flag]:
    return [
        _Flag(
            flag_id="SUSPICIOUS_IMPROVEMENT_sub1_strategic_saturation_score",
            flag_type="SUSPICIOUS_IMPROVEMENT",
            dimension="Strategic Saturation",
            description="Strategic Saturation improved by 18.0 points (72 → 54).",
            detected_at="2026-06-01T00:00:00+00:00",
        ),
        _Flag(
            flag_id="DATA_GAP_CLOSURE_sub1_reporting_integrity_score",
            flag_type="DATA_GAP_CLOSURE",
            dimension="Reporting Integrity",
            description="Reporting Integrity prior score (75) was near missing-data default.",
            detected_at="2026-06-01T00:00:00+00:00",
        ),
    ]


# ------------------------------------------------------------------
# Tests
# ------------------------------------------------------------------

def test_output_path_naming(source_workbook, two_flags):
    output = write_variance_flags(source_workbook, two_flags, "Alpha Corp", "2026-Q2")
    assert output.parent == source_workbook.parent
    assert output.name == "client_alpha_CPOI_flagged.xlsx"
    assert output.exists()


def test_source_not_modified(source_workbook, two_flags):
    mtime_before = source_workbook.stat().st_mtime
    write_variance_flags(source_workbook, two_flags, "Alpha Corp", "2026-Q2")
    mtime_after = source_workbook.stat().st_mtime
    assert mtime_before == mtime_after

    # Original should not have the VARIANCE_FLAGS sheet
    wb_orig = load_workbook(source_workbook)
    assert _SHEET_NAME not in wb_orig.sheetnames


def test_sheet_created(source_workbook, two_flags):
    output = write_variance_flags(source_workbook, two_flags, "Alpha Corp", "2026-Q2")
    wb = load_workbook(output)
    assert _SHEET_NAME in wb.sheetnames


def test_header_row(source_workbook, two_flags):
    output = write_variance_flags(source_workbook, two_flags, "Alpha Corp", "2026-Q2")
    wb = load_workbook(output)
    ws = wb[_SHEET_NAME]
    # Header is row 3 (row 1 = title, row 2 = blank)
    header_values = [ws.cell(row=3, column=i).value for i in range(1, len(_HEADERS) + 1)]
    assert header_values == _HEADERS


def test_flag_rows_written(source_workbook, two_flags):
    output = write_variance_flags(source_workbook, two_flags, "Alpha Corp", "2026-Q2")
    wb = load_workbook(output)
    ws = wb[_SHEET_NAME]

    # Data starts at row 4 (after title, blank, header)
    row4 = [ws.cell(row=4, column=i).value for i in range(1, len(_HEADERS) + 1)]
    row5 = [ws.cell(row=5, column=i).value for i in range(1, len(_HEADERS) + 1)]

    flag_types = {row4[_HEADERS.index("Flag_Type")], row5[_HEADERS.index("Flag_Type")]}
    assert flag_types == {"SUSPICIOUS_IMPROVEMENT", "DATA_GAP_CLOSURE"}


def test_flag_severity_values(source_workbook, two_flags):
    output = write_variance_flags(source_workbook, two_flags, "Alpha Corp", "2026-Q2")
    wb = load_workbook(output)
    ws = wb[_SHEET_NAME]

    sev_col = _HEADERS.index("Severity") + 1
    severities = {ws.cell(row=r, column=sev_col).value for r in [4, 5]}
    assert "CRITICAL" in severities   # SUSPICIOUS_IMPROVEMENT
    assert "ELEVATED" in severities   # DATA_GAP_CLOSURE


def test_no_flags_message(source_workbook):
    output = write_variance_flags(source_workbook, [], "Alpha Corp", "2026-Q2")
    wb = load_workbook(output)
    ws = wb[_SHEET_NAME]
    cell_value = ws.cell(row=4, column=1).value
    assert cell_value is not None
    assert "No anomaly flags" in cell_value


def test_file_not_found_raises(tmp_path):
    missing = tmp_path / "does_not_exist.xlsx"
    with pytest.raises(FileNotFoundError, match="does_not_exist.xlsx"):
        write_variance_flags(missing, [], "Nobody", "2026-Q1")


def test_existing_sheet_replaced(source_workbook, two_flags):
    # Write once, then write again with different flags — sheet should reflect second write
    write_variance_flags(source_workbook, two_flags, "Alpha Corp", "2026-Q2")

    output = source_workbook.parent / "client_alpha_CPOI_flagged.xlsx"
    single_flag = [two_flags[0]]
    write_variance_flags(source_workbook, single_flag, "Alpha Corp", "2026-Q2")

    wb = load_workbook(output)
    ws = wb[_SHEET_NAME]
    # Row 5 should be empty (only 1 flag written)
    assert ws.cell(row=5, column=1).value is None
