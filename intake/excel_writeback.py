"""
intake/excel_writeback.py
CPOI Platform -- Outbound Excel write-back (Module 7).

Adds a VARIANCE_FLAGS tab to a copy of the client's workbook.
The original workbook is never modified; output is written to a
sibling file with a '_CPOI_flagged' suffix.

Usage:
    from intake.excel_writeback import write_variance_flags
    output_path = write_variance_flags(source_path, flags, client_name, period_label)
"""

from __future__ import annotations

import shutil
from pathlib import Path
from typing import TYPE_CHECKING

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

if TYPE_CHECKING:
    from scoring.anomaly import AnomalyFlag

# ------------------------------------------------------------------
# Constants
# ------------------------------------------------------------------

_SHEET_NAME = "VARIANCE_FLAGS"

_HEADERS = [
    "Tab_Source",
    "Row_Reference",
    "Flag_Type",
    "Platform_Finding",
    "Recommended_Client_Review",
    "Severity",
]

_RECOMMENDED_REVIEW: dict[str, str] = {
    "SUSPICIOUS_IMPROVEMENT": (
        "Review source data for this dimension. Confirm that the improvement "
        "reflects genuine operational changes and is not the result of "
        "revised assumptions or re-categorized data."
    ),
    "INCONSISTENT_SIGNAL": (
        "Cross-check the primary signal reading against the dimension score. "
        "If the signal reading is accurate, re-examine sub-category inputs to "
        "identify the source of the divergence."
    ),
    "DATA_GAP_CLOSURE": (
        "Confirm that previously missing data was provided from a legitimate "
        "source this period. Document the source of the newly available data "
        "in the submission notes."
    ),
    "ZERO_MOVEMENT_STAGNATION": (
        "Verify that the client submitted updated operational data for this "
        "period. If data is unchanged, note whether that reflects genuine "
        "operational stability or a failure to update the submission."
    ),
}

_FLAG_SEVERITY: dict[str, str] = {
    "SUSPICIOUS_IMPROVEMENT":  "CRITICAL",
    "INCONSISTENT_SIGNAL":     "CRITICAL",
    "DATA_GAP_CLOSURE":        "ELEVATED",
    "ZERO_MOVEMENT_STAGNATION": "WATCH",
}

# openpyxl PatternFill ARGB values (FF = fully opaque)
_SEVERITY_FILL: dict[str, PatternFill] = {
    "CRITICAL": PatternFill(start_color="FFFCE4E4", end_color="FFFCE4E4", fill_type="solid"),
    "ELEVATED": PatternFill(start_color="FFFFF3E0", end_color="FFFFF3E0", fill_type="solid"),
    "WATCH":    PatternFill(start_color="FFFFFDE7", end_color="FFFFFDE7", fill_type="solid"),
}

_HEADER_FILL = PatternFill(start_color="FFE0E0E0", end_color="FFE0E0E0", fill_type="solid")
_HEADER_FONT = Font(bold=True)

# Approximate column widths (characters)
_COL_WIDTHS = {
    "Tab_Source":                  24,
    "Row_Reference":               42,
    "Flag_Type":                   28,
    "Platform_Finding":            70,
    "Recommended_Client_Review":   70,
    "Severity":                    12,
}


# ------------------------------------------------------------------
# Public API
# ------------------------------------------------------------------

def write_variance_flags(
    source_path: Path,
    flags: list["AnomalyFlag"],
    client_name: str,
    period_label: str,
) -> Path:
    """
    Write a VARIANCE_FLAGS sheet to a copy of the client's workbook.

    Args:
        source_path:  Path to the original client Excel workbook.
        flags:        List of AnomalyFlag objects from detect_anomaly_flags().
        client_name:  Used only in the sheet title comment row.
        period_label: Human-readable period string (e.g. "2026-Q2").

    Returns:
        Path to the output workbook (sibling of source_path with
        '_CPOI_flagged' suffix).

    Raises:
        FileNotFoundError: if source_path does not exist.
    """
    if not source_path.exists():
        raise FileNotFoundError(
            f"Source workbook not found: {source_path}. "
            "Ensure the client workbook path is correct before calling write_variance_flags()."
        )

    output_path = source_path.parent / f"{source_path.stem}_CPOI_flagged{source_path.suffix}"
    shutil.copy2(source_path, output_path)

    wb = load_workbook(output_path)

    # Remove existing sheet if present so we can rebuild it cleanly
    if _SHEET_NAME in wb.sheetnames:
        del wb[_SHEET_NAME]

    ws = wb.create_sheet(title=_SHEET_NAME)

    # ------------------------------------------------------------------
    # Title row
    # ------------------------------------------------------------------
    ws.append([f"CPOI Variance Flags -- {client_name} -- Period: {period_label}"])
    title_cell = ws.cell(row=1, column=1)
    title_cell.font = Font(bold=True, size=12)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(_HEADERS))
    ws.append([])  # blank separator row

    # ------------------------------------------------------------------
    # Header row
    # ------------------------------------------------------------------
    header_row_idx = 3
    ws.append(_HEADERS)
    for col_idx, header in enumerate(_HEADERS, start=1):
        cell = ws.cell(row=header_row_idx, column=col_idx)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    # ------------------------------------------------------------------
    # Data rows
    # ------------------------------------------------------------------
    if not flags:
        no_flags_row = header_row_idx + 1
        ws.cell(row=no_flags_row, column=1, value="No anomaly flags detected for this period.")
        ws.cell(row=no_flags_row, column=1).font = Font(italic=True)
    else:
        for flag in flags:
            severity = _FLAG_SEVERITY.get(flag.flag_type, "WATCH")
            tab_source = flag.dimension if flag.dimension else "Portfolio"
            recommended = _RECOMMENDED_REVIEW.get(flag.flag_type, "Review flag details.")

            row_data = [
                tab_source,
                flag.flag_id,
                flag.flag_type,
                flag.description,
                recommended,
                severity,
            ]
            ws.append(row_data)

            data_row_idx = ws.max_row
            sev_fill = _SEVERITY_FILL.get(severity)
            for col_idx in range(1, len(_HEADERS) + 1):
                cell = ws.cell(row=data_row_idx, column=col_idx)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if sev_fill and col_idx == len(_HEADERS):
                    cell.fill = sev_fill

    # ------------------------------------------------------------------
    # Column widths
    # ------------------------------------------------------------------
    for col_idx, header in enumerate(_HEADERS, start=1):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = _COL_WIDTHS.get(header, 20)

    # Row height for data rows (wrap_text needs explicit height to render)
    for row_idx in range(header_row_idx + 1, ws.max_row + 1):
        ws.row_dimensions[row_idx].height = 60

    wb.save(output_path)
    return output_path
