"""
intake/template.py
CPOI Platform -- Blank intake workbook generator.

Builds the standardized CP_OEI_DataIntake workbook (the 8-tab compliance
contract) with the exact headers the validator expects. Headers are derived
directly from intake.validator.SCHEMA so the template can never drift from
validation rules. The client fills it in; the platform reads it back.

Usage:
    from intake.template import build_blank_template, template_filename
    data = build_blank_template()          # -> bytes (.xlsx)
"""

from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from intake.validator import (
    REQUIRED_TABS, SCHEMA, TAB_METADATA, METADATA_REQUIRED_FIELDS,
)

# Light guidance shown in the METADATA value cells (cleared by the client).
_METADATA_HINTS = {
    "Reporting_Period_Start": "YYYY-MM-DD",
    "Reporting_Period_End": "YYYY-MM-DD",
    "Submission_Date": "YYYY-MM-DD",
    "Engagement_Type": "Snapshot or OEIL_Month_1",
    "Data_Version": "1.0",
}

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill(start_color="2E4A6A", end_color="2E4A6A", fill_type="solid")
_KEY_FONT = Font(bold=True)


def build_blank_template() -> bytes:
    """Return the blank intake workbook as .xlsx bytes."""
    wb = Workbook()
    wb.remove(wb.active)  # drop the default sheet

    for tab in REQUIRED_TABS:
        ws = wb.create_sheet(title=tab)

        if tab == TAB_METADATA:
            for field in METADATA_REQUIRED_FIELDS:
                ws.append([field, _METADATA_HINTS.get(field, "")])
            for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=1):
                row[0].font = _KEY_FONT
            ws.column_dimensions["A"].width = 26
            ws.column_dimensions["B"].width = 32
            continue

        schema = SCHEMA[tab]
        headers = list(schema.required_columns) + list(schema.optional_columns)
        ws.append(headers)
        for cell in ws[1]:
            cell.font = _HEADER_FONT
            cell.fill = _HEADER_FILL
        for idx in range(1, len(headers) + 1):
            ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = 22

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def template_filename() -> str:
    """A friendly default download name including the current month."""
    return f"CP_OEI_DataIntake_Template_{date.today():%Y_%m}.xlsx"
